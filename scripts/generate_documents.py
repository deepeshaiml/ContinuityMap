#!/usr/bin/env python3
"""Generate ContinuityMap idea, business, evaluation, and pitch PDF artifacts."""

from __future__ import annotations

import json
from datetime import date
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch, mm
from reportlab.platypus import (
    KeepTogether,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parents[1]
REPORTS = ROOT / "reports"
REPORTS.mkdir(exist_ok=True)
NAVY = colors.HexColor("#102638")
INK = colors.HexColor("#203747")
TEAL = colors.HexColor("#087F8C")
MUTED = colors.HexColor("#71858D")
LINE = colors.HexColor("#DFE8EA")
PALE = colors.HexColor("#F2F7F7")
AMBER = colors.HexColor("#B17B24")
CORAL = colors.HexColor("#C7543F")


PAGES = [
    {
        "kicker": "01 · PRODUCT THESIS",
        "title": "ContinuityMap at a glance",
        "dek": "A transparent planning workspace for understanding declared dependencies between essential services.",
        "paras": [
            "Electricity, water, hospitals, communications, transport, and emergency response rely on other services and resources. The information is often spread across continuity plans, equipment lists, service descriptions, and exercise notes. During an exercise, teams may have to piece together who relies on whom, whether a backup is confirmed, and where each claim came from.",
            "ContinuityMap puts those recorded relationships in one map. A planner can select a service, see what it relies on and what relies on it, check the evidence, and run the same what-if exercise again. Only relationships a human has marked verified affect the calculated result. Suggested links and open questions stay separate.",
            "The result is a planning aid, not a forecast. It follows the recorded links and stated backup rules. It does not predict physical damage, actual power or water flow, how long a backup lasts, when service returns, or whether an organization is ready. If those facts are missing, the system leaves them unknown.",
            "Example: when the fictional Northbank Substation Alpha is disrupted, water treatment remains available in the exercise because a verified rule says either Alpha or the Civic Solar Microgrid can supply it. A hospital is marked unavailable because its required health-records path runs through a data center powered by Alpha. This is what the demo’s records say; it is not a prediction about a real hospital.",
        ],
        "bullets": [
            "A fictional Northbank network with 43 items, 80 relationship records, and three repeatable exercises.",
            "Human approval is required before a proposed relationship can affect an official scenario result.",
            "The demo can run locally. Developers can use PostgreSQL, FastAPI, React, NetworkX, and optional Ollama.",
        ],
        "callout": "In the Alpha exercise, 17 entries are marked unavailable (the substation plus 16 downstream entries); two open questions are shown for review, not counted as failures.",
    },
    {
        "kicker": "02 · PROBLEM & RESEARCH",
        "title": "A coordination problem with real policy context",
        "dek": "The challenge is maintaining a shared view of dependencies and evidence across organizational boundaries.",
        "paras": [
            "Resilience planning involves service owners, public authorities, utilities, hospitals, transport teams, suppliers, and emergency groups. One team may own a plan, another a facility, and another the communications service everyone uses. A useful shared map therefore needs more than lines: it needs a clear arrow direction, a named owner, supporting evidence, and a date when someone last checked the claim.",
            "Public guidance points to the same planning need. NATO’s 2026 resilience baseline discusses dependencies and cascading effects across essential sectors. FEMA asks planners to consider which functions support or depend on other functions. CISA/FEMA cyber-incident guidance asks planners to look upstream and downstream. NIST SP 800-34 is useful for information-system recovery, though it does not cover the whole cross-sector problem.",
            "A 2024 research review found that infrastructure studies use different dependency labels and can lose track of which way a relationship points. ContinuityMap makes the direction visible: provider → service that depends on it. This is a practical design choice, not a claim that one universal research standard exists.",
            "Existing resilience platforms already offer maps, dashboards, and dependency features. For example, Everbridge publicly describes Bridge tools for mapping services and dependencies, finding gaps, setting impact tolerances, and linking plans to response workflows. ContinuityMap should not claim those basic features are new. Its early idea is to make each relationship and its review status easy to inspect, then show repeatable results from only owner-confirmed links. Buyer comparisons are still needed.",
        ],
        "bullets": [
            "Likely initial users: resilience offices, utilities, hospitals, airports, infrastructure consultancies, and emergency planning teams.",
            "Likely customer problem: cross-team preparation and verification of dependency-aware exercises.",
            "Primary limitation: accurate, current dependency data is customer-specific and often hard to assemble.",
        ],
        "callout": "Research links, source limits, and software license notes are collected in docs/SOURCES.md. Competitor fit and customer demand still need interviews and product comparisons.",
    },
    {
        "kicker": "03 · PRODUCT & WORKFLOW",
        "title": "From inventory to owner-reviewed action",
        "dek": "The core workflow keeps the graph and its evidence in the same place.",
        "paras": [
            "The inventory lists services, facilities, equipment, and backup resources. Every relationship records what supplies what, the type of support, a source quote and reference, its review state, an owner, any stated conditions, and the date it was last checked. A planner can select an item on the map, inspect its relationships and evidence, filter by review state, and edit a record.",
            "There are four review states. Verified means a person approved the claim and it can affect an exercise. Proposed means a possible link is waiting for review. Unknown means the source leaves a question open. Rejected means a reviewer declined the claim. Proposed, unknown, and rejected links do not make a service unavailable in the calculation.",
            "A scenario starts with one or more resources chosen as unavailable. The app follows verified relationships, shows each result and its evidence path, and lists open questions separately. Users can edit the graph, rerun the same scenario, and export an impact report, a gap report, or a human-review checklist.",
            "The Northbank demo covers electricity, water, hospitals, clinics, communications, emergency operations, transport, shelters, data services, fuel, and backup resources. All names and relationships are invented; it is not a map of a real city or national infrastructure.",
        ],
        "bullets": [
            "Load a JSON file and check for missing or conflicting relationship details.",
            "Select a scenario or enter a custom disruption set.",
            "Review results, open questions, and owner actions; export PDF reports.",
        ],
        "callout": "A scenario can only be as complete as the relationship records people have collected, checked, and kept current.",
    },
    {
        "kicker": "04 · DATA MODEL & SIMULATION",
        "title": "A simple graph with explicit redundancy rules",
        "dek": "Only verified, directed relationships determine official impact status.",
        "paras": [
            "Each arrow points from a provider to the service that relies on it. A dependent service can have one or more groups of required providers. ANY means at least one provider in that group is enough—for example, utility power or a generator. ALL means every provider is needed—for example, power and network access. When a service has multiple groups, each group must be satisfied.",
            "The calculation starts with the resource or resources selected as disrupted. It checks verified groups and adds a service if one of its required groups is no longer met. Then it checks again, because newly unavailable services may support other services. It stops when another full check adds nothing. Cycles are flagged for review, but the repeated check cannot run forever.",
            "Proposed and unknown relationships are handled as review questions. If their provider is disrupted, the app may flag the dependent for an owner to check, but it does not mark it unavailable because of that unverified link. Rejected links are excluded. The validator also checks missing nodes, duplicate or conflicting relationships, unsupported assumptions, and missing review dates.",
            "The engine does not model capacity, flow, load shedding, time-to-failure, restoration time, physical damage, shared causes, or probability. If the owner has not supplied a value, it stays unknown. The result means only that a declared requirement is or is not met in this exercise.",
        ],
        "bullets": [
            "Substation Alpha unavailable: 17 unavailable entries, including the starting substation.",
            "East water pumping loses utility and backup power: 20 unavailable nodes.",
            "Regional communications outage: 9 unavailable nodes; independent radio remains available.",
        ],
            "callout": "In the Alpha exercise, 2 possible paths are shown for review. They are not counted as unavailable services.",
    },
    {
        "kicker": "05 · AI & GOVERNANCE",
        "title": "AI proposes; owners decide",
        "dek": "Local document extraction supports review without becoming an authority over the graph.",
        "paras": [
            "A user can upload an invented TXT, Markdown, CSV, PDF, or DOCX plan. The app looks for names already in the graph and asks an optional local Ollama model to suggest who depends on whom, with a supporting quote. The names must match existing graph items, and the quote must appear in the uploaded text. If a model is not configured or its output is unusable, the app uses a clearly labeled rule-based method.",
            "Every extracted link starts as proposed and keeps its quote and document reference. A service owner can correct the link, add missing details, verify it, or reject it. Until a person verifies it, the link cannot change an official exercise result.",
            "The continuity assistant starts from the calculated result, relationship evidence, open questions, recovery resources, and written exercise actions. It may draft a summary and suggested actions, but the user must review them. Citations must point to known evidence references. If the model response cannot be checked, the app uses a structured evidence summary instead. The assistant does not change service status or a recovery plan.",
            "These checks reduce some AI risks but do not guarantee that an AI summary is correct. A valid citation ID does not prove that a sentence interprets the source correctly. Suggestions need human review. Five simple synthetic extraction examples are only a basic software check, not proof the tool understands real plans.",
        ],
        "bullets": [
            "2 MB upload limit and 80,000 extracted character limit in the prototype.",
            "No cloud LLM is configured or required.",
            "An assessment is not a readiness score or an approved recovery plan.",
        ],
        "callout": "In plain terms: AI can point to a possible link; an accountable service owner decides whether it belongs in the verified graph.",
    },
    {
        "kicker": "06 · ARCHITECTURE & DEPLOYMENT",
        "title": "A conventional, self-hostable application stack",
        "dek": "The MVP deliberately keeps the analysis core independent from web, persistence, and AI code.",
        "paras": [
            "The interface uses React and TypeScript; Cytoscape draws and lets users edit the network. FastAPI and Pydantic provide and check the API fields. NetworkX checks and calculates the graph. SQLAlchemy saves records in PostgreSQL. Local development and tests can use SQLite. Docker Compose starts the web app, API, and database. Ollama is optional and can run a language model on the same machine or local network.",
            "For the current 43-node demo, PostgreSQL plus NetworkX is simpler than adding a separate graph database. Larger deployments, multiple customer groups, graph history, or time-based records may change that choice; those needs should be confirmed with users first.",
            "Uploaded source documents are read in memory and not retained as files. Extracted quotes and references may be saved with a proposed relationship. This development API has no user login or permissions; it accepts only explicitly fictional nodes. Do not expose it publicly or use real infrastructure data with it.",
            "A customer-ready system would need user accounts, permission levels for editing and approving, separation between customers, secret handling, encrypted connections, backups, data-retention rules, upload scanning, audit monitoring, recovery plans, and a security review. An on-premises option may help a customer control data but increases setup and support work.",
        ],
        "bullets": [
            "PostgreSQL saves the demo data; SQLite is available for local development.",
            "No cloud AI is required. A local Ollama model is optional.",
            "Pinned software versions, license notes, compatibility, and security checks are listed in docs/SOURCES.md.",
        ],
        "callout": "See ContinuityMap_Architecture.svg for the picture and ContinuityMap_Technical_Design.md for developer details and API routes.",
    },
    {
        "kicker": "07 · EVALUATION & IMPACT",
        "title": "Measure the implementation, then test with real owners",
        "dek": "Current evaluation checks deterministic software behavior against fixed synthetic fixtures.",
        "paras": [
            "The automated checks use small example networks and the 43-item fictional graph. They check backup rules, cycles, disconnected services, missing references, simultaneous disruptions, open questions, conflicting claims, document suggestions, imports, API behavior, and PDF reports. The evaluation also asks whether the app extracts simple examples, keeps human review states straight, uses verified relationships only, preserves uncertainty, points to existing evidence IDs, and includes required report sections.",
            "The latest snapshot shows 5 of 5 simple extraction examples matched, all 3 scenario result lists matched their separately prepared test answers, no unverified relationship caused a service to be marked unavailable, and the required report sections were present. Precision means “of the links suggested, how many were expected?” Recall means “of the expected links, how many were found?” These results describe a few fixed synthetic examples, not real customer documents or real-world disruption accuracy.",
            "Possible customer benefits include less repeated manual scenario preparation, quicker evidence review, earlier visibility of missing backup documentation, better coordination, and clearer records of who reviewed what. These are ideas to test, not proven savings. A pilot could measure preparation time, how many important relationships have current owners and evidence, whether owners agree with the paths, and whether follow-up actions are completed.",
            "The software cannot repair an incomplete or disputed graph on its own. A customer would need to map their own terms and records, assign accountable owners, decide what counts as verified, and schedule regular review.",
        ],
        "bullets": [
            "Automated suite snapshot: 29 tests passed in the recorded implementation run.",
            "Graph fixture: 43 nodes, 80 edge records, 78 verified edges, 2 explicit unknown questions.",
            "Median local calculation time: about 0.04 ms in one test environment; hardware-specific and not a production benchmark.",
        ],
        "callout": "The scores show that the software matches its small, hand-checked test cases. They do not show that the fictional graph is complete or that a real service would fail.",
    },
    {
        "kicker": "08 · BUSINESS CASE, RISKS & SOURCES",
        "title": "Commercial path and responsible next steps",
        "dek": "Implementation quality and customer data readiness are central to both value and gross margin.",
        "paras": [
            "Here is one illustrative pricing example, not a quote or forecast. Charge $65,000 per year for the organization plus $9,000 per site for three sites: $65,000 + (3 × $9,000) = $92,000 in recurring annual revenue. Add $28,000 for initial onboarding, so first-year revenue is $120,000. Assume direct delivery costs of $18,000 for onboarding plus $12,000 for one year of cloud and support: $30,000 total. First-year gross profit is $120,000 − $30,000 = $90,000, or 75% of revenue. Gross profit leaves out research and development, sales, management, tax, security assurance, and other company costs. If year two costs only $12,000 to host and support, recurring gross margin would be 87%; this too is an assumption.",
            "Possible ways to charge include an organization subscription, a fee per site, an annual customer-hosted license, data-integration work, optional scenario modules, and professional dependency assessments. A broad estimate of time spent preparing exercises suggests $7,200 to $134,400 in possible annual staff-time value, depending on hours, exercise frequency, and staff cost. None of those inputs has been checked with buyers. At the illustrative middle case ($32,000), staff-time savings alone are less than the assumed $92,000 annual subscription. The product may need to show additional value from multi-site coordination, evidence traceability, audits, or consultant time.",
            "Main risks include incomplete or disputed dependency data, sensitive records, AI mistakes, misunderstood results, stale owners, long procurement, and costly customer-specific setup. Start with a narrow approved dataset, named owners, one reviewed exercise, and a controlled deployment. Do not claim to predict outages or certify compliance.",
            "Possible roadmap: months 0–3, interview planners and confirm a safe use case; months 3–6, run a customer-controlled pilot with human-reviewed data and basic identity controls; months 6–12, add approved integrations and stronger access, audit, and security features. Continue only if owners trust the evidence paths and setup effort can be repeated.",
        ],
        "bullets": [
            "Data completeness and integration effort are the commercial gating factors.",
            "Everbridge and other established resilience platforms create expected feature overlap; validate fit in buyer-led comparisons.",
            "The best first pilot is small, has named owners, and uses evidence the team can review.",
        ],
        "callout": "All prices, costs, and labor values are assumptions for discussion. This is a product concept, not a financial forecast or a legal, safety, or compliance certification.",
    },
]


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, value, end])


def build_idea_docx() -> None:
    document = Document()
    section = document.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.28)
    section.footer_distance = Inches(0.28)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(43, 64, 75)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.12
    for name, size, color in [("Heading 1", 22, "102638"), ("Heading 2", 11, "087F8C")]:
        style = document.styles[name]
        style.font.name = "Aptos Display" if name == "Heading 1" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(2)
        style.paragraph_format.space_after = Pt(5)
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("CONTINUITYMAP   ·   IDEA DOCUMENT")
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string("087F8C")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("SYNTHETIC PLANNING CONCEPT     ·     ")
    run.font.size = Pt(7)
    run.font.color.rgb = RGBColor.from_string("71858D")
    add_page_number(footer)

    for index, page in enumerate(PAGES):
        kicker = document.add_paragraph()
        kicker.paragraph_format.space_after = Pt(4)
        run = kicker.add_run(page["kicker"])
        run.bold = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string("087F8C")
        title = document.add_paragraph(style="Heading 1")
        title.add_run(page["title"])
        dek = document.add_paragraph()
        dek.paragraph_format.space_after = Pt(10)
        run = dek.add_run(page["dek"])
        run.italic = True
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor.from_string("617983")
        for paragraph_text in page["paras"]:
            paragraph = document.add_paragraph(paragraph_text)
            paragraph.paragraph_format.keep_together = True
        for bullet in page["bullets"]:
            paragraph = document.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.left_indent = Inches(0.22)
            paragraph.paragraph_format.first_line_indent = Inches(-0.12)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.add_run(bullet)
        table = document.add_table(rows=1, cols=1)
        table.autofit = True
        cell = table.cell(0, 0)
        set_cell_shading(cell, "EAF4F2")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        callout = cell.paragraphs[0]
        callout.paragraph_format.space_before = Pt(4)
        callout.paragraph_format.space_after = Pt(4)
        run = callout.add_run(page["callout"])
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor.from_string("275A5C")
        if index < len(PAGES) - 1:
            document.add_page_break()
    document.core_properties.title = "ContinuityMap Idea Document"
    document.core_properties.subject = "AI-assisted infrastructure dependency and resilience planning"
    document.core_properties.author = "ContinuityMap"
    document.save(ROOT / "ContinuityMap_Idea_Document.docx")


def base_styles():
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="CMTitle", parent=styles["Title"], fontName="Helvetica-Bold", fontSize=22, leading=26, textColor=NAVY, spaceAfter=6, alignment=TA_LEFT))
    styles.add(ParagraphStyle(name="CMSub", parent=styles["Normal"], fontName="Helvetica-Oblique", fontSize=9.5, leading=13, textColor=MUTED, spaceAfter=9))
    styles.add(ParagraphStyle(name="CMBody", parent=styles["BodyText"], fontName="Helvetica", fontSize=9.2, leading=13, textColor=INK, spaceAfter=6))
    styles.add(ParagraphStyle(name="CMKicker", parent=styles["Normal"], fontName="Helvetica-Bold", fontSize=7.5, leading=9, textColor=TEAL, spaceAfter=4, charSpace=0.7))
    styles.add(ParagraphStyle(name="CMHeading", parent=styles["Heading2"], fontName="Helvetica-Bold", fontSize=11, leading=14, textColor=TEAL, spaceBefore=5, spaceAfter=5))
    styles.add(ParagraphStyle(name="CMBullet", parent=styles["BodyText"], fontName="Helvetica", fontSize=9, leading=12.2, textColor=INK, leftIndent=12, firstLineIndent=-8, bulletIndent=0, spaceAfter=3))
    styles.add(ParagraphStyle(name="CMSmall", parent=styles["Normal"], fontName="Helvetica", fontSize=8, leading=10.2, textColor=MUTED, spaceAfter=3))
    styles.add(ParagraphStyle(name="CMTable", parent=styles["BodyText"], fontName="Helvetica", fontSize=7.8, leading=9.6, textColor=INK, spaceAfter=0))
    styles.add(ParagraphStyle(name="CMTableHead", parent=styles["BodyText"], fontName="Helvetica-Bold", fontSize=7.8, leading=9.6, textColor=colors.white, spaceAfter=0))
    return styles


def pdf_header_footer(canvas, doc):
    canvas.saveState()
    w, h = doc.pagesize
    canvas.setFillColor(NAVY)
    canvas.rect(0, h - 7 * mm, w, 7 * mm, fill=1, stroke=0)
    canvas.setFillColor(TEAL)
    canvas.rect(0, h - 7 * mm, 5 * mm, 7 * mm, fill=1, stroke=0)
    canvas.setFont("Helvetica-Bold", 7)
    canvas.setFillColor(colors.HexColor("#E8F1F3"))
    canvas.drawString(17 * mm, h - 5 * mm, "CONTINUITYMAP  ·  SYNTHETIC PLANNING CONCEPT")
    canvas.setStrokeColor(LINE)
    canvas.line(17 * mm, 12 * mm, w - 17 * mm, 12 * mm)
    canvas.setFillColor(MUTED)
    canvas.setFont("Helvetica", 7)
    canvas.drawString(17 * mm, 7 * mm, "Decision support only · Verify with accountable service owners")
    canvas.drawRightString(w - 17 * mm, 7 * mm, str(doc.page))
    canvas.restoreState()


def build_idea_pdf() -> None:
    styles = base_styles()
    story = []
    for index, page in enumerate(PAGES):
        story.append(Paragraph(page["kicker"], styles["CMKicker"]))
        story.append(Paragraph(escape(page["title"]), styles["CMTitle"]))
        story.append(Paragraph(escape(page["dek"]), styles["CMSub"]))
        for paragraph_text in page["paras"]:
            story.append(Paragraph(escape(paragraph_text), styles["CMBody"]))
        for bullet in page["bullets"]:
            story.append(Paragraph("• " + escape(bullet), styles["CMBullet"]))
        callout = Table([[Paragraph(escape(page["callout"]), styles["CMBody"])]], colWidths=[170 * mm])
        callout.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#EAF4F2")), ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#CBE3DE")), ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 8), ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
        story.extend([Spacer(1, 4), callout])
        if index < len(PAGES) - 1:
            story.append(PageBreak())
    document = SimpleDocTemplate(str(ROOT / "ContinuityMap_Idea_Document.pdf"), pagesize=A4, rightMargin=17 * mm, leftMargin=17 * mm, topMargin=17 * mm, bottomMargin=18 * mm, title="ContinuityMap Idea Document", author="ContinuityMap")
    document.build(story, onFirstPage=pdf_header_footer, onLaterPages=pdf_header_footer)


def table_data(rows, styles, widths=None):
    converted = []
    for row_index, row in enumerate(rows):
        style = styles["CMTableHead"] if row_index == 0 else styles["CMTable"]
        converted.append([Paragraph(escape(str(cell)), style) for cell in row])
    table = Table(converted, colWidths=widths, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("GRID", (0, 0), (-1, -1), 0.35, LINE),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, PALE]),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4), ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def build_business_pdf() -> None:
    styles = base_styles()
    story = [Paragraph("BUSINESS CASE · ILLUSTRATIVE ASSUMPTIONS", styles["CMKicker"]), Paragraph("A focused planning product with an implementation-led path", styles["CMTitle"]), Paragraph("Commercial hypotheses for validation with resilience teams; not a forecast, offer, or customer quote.", styles["CMSub"])]
    story.append(Paragraph("Customer value hypothesis", styles["CMHeading"]))
    story.append(Paragraph("ContinuityMap could reduce repeated manual assembly of dependency-aware exercises, make evidence gaps visible earlier, improve coordination between service owners, and produce traceable review records. These benefits are not quantified by the synthetic application. A pilot should measure preparation time, graph coverage, owner agreement with impact paths, and follow-up action closure.", styles["CMBody"]))
    story.append(Paragraph("Illustrative delivery and pricing models", styles["CMHeading"]))
    story.append(Paragraph("Possible offers include an annual organization subscription, a fee for each site, a customer-hosted license, setup and data-mapping work, optional scenario modules, and professional assessments. This example assumes a $65,000 annual platform fee plus $9,000 per site for three sites, or $92,000 recurring per organization. First-year revenue adds $28,000 onboarding.", styles["CMBody"]))
    rows = [["Organizations", "Year-one revenue", "Direct cost", "Gross profit", "Gross margin", "Annual subscription total"], ["1", "$120,000", "$30,000", "$90,000", "75%", "$92,000"], ["5", "$600,000", "$150,000", "$450,000", "75%", "$460,000"], ["10", "$1,200,000", "$300,000", "$900,000", "75%", "$920,000"], ["25", "$3,000,000", "$750,000", "$2,250,000", "75%", "$2,300,000"]]
    story.append(table_data(rows, styles, [25 * mm, 29 * mm, 24 * mm, 27 * mm, 22 * mm, 30 * mm]))
    story.append(Spacer(1, 4))
    story.append(Paragraph("How to read this example: direct cost is the assumed cost to onboard and serve an organization. Gross profit = revenue − direct costs. Gross margin = gross profit ÷ revenue. For one organization, $120,000 − $30,000 = $90,000 gross profit; $90,000 ÷ $120,000 = 75%. The $30,000 assumes $18,000 onboarding delivery plus $12,000 yearly hosting/support. Year two would have $80,000 gross profit and 87% gross margin only if revenue stays $92,000 and direct cost stays $12,000. These figures exclude product development, sales, management, taxes, security assurance, and other company overhead. They are assumptions, not a forecast.", styles["CMSmall"]))
    story.append(PageBreak())
    story.extend([Paragraph("VALUE SENSITIVITY · NOT A PROMISE", styles["CMKicker"]), Paragraph("Labor savings alone may not support the entry price", styles["CMTitle"]), Paragraph("A transparent sensitivity range exposes the willingness-to-pay question rather than hiding it.", styles["CMSub"])])
    sensitivity = [["Example", "Staff hours saved / exercise", "Exercises / year", "Staff cost / hour", "Possible yearly staff-time value"], ["Low", "40", "3", "$60", "$7,200"], ["Middle", "80", "4", "$100", "$32,000"], ["High", "160", "6", "$140", "$134,400"]]
    story.append(table_data(sensitivity, styles, [33 * mm, 29 * mm, 31 * mm, 30 * mm, 34 * mm]))
    story.append(Spacer(1, 8))
    for paragraph in [
        "Formula: staff hours saved per exercise × exercises per year × staff cost per hour. No target customer has validated these inputs. The middle example is less than the assumed $92,000 annual subscription, so staff-time savings alone may not justify the price for a single-site or low-frequency buyer.",
        "The commercial case may depend on multi-site rollups, reduced consultant effort, audit traceability, repeatable exercises, service-owner coordination, and lower cost of updating plans. Any claimed cash value requires customer baseline data and a controlled before/after study. The product cannot guarantee avoided outages or improved regulatory status.",
    ]:
        story.append(Paragraph(paragraph, styles["CMBody"]))
    story.append(Paragraph("Customer segments and purchase constraints", styles["CMHeading"]))
    rows = [["Segment", "Potential job", "Adoption issue"], ["Municipal resilience office", "Cross-agency exercise preparation", "Data-sharing authority and owner coordination"], ["Power / water utility", "Facility and service dependency review", "Sensitive operational data and system integration"], ["Hospital / airport", "Continuity scenario traceability", "Procurement, identity, and existing BCM stack"], ["Consultancy / government", "Repeatable assessment and on-prem delivery", "Custom data onboarding and assurance burden"]]
    story.append(table_data(rows, styles, [38 * mm, 60 * mm, 59 * mm]))
    story.append(PageBreak())
    story.extend([Paragraph("COSTS · RISKS · GO-TO-MARKET", styles["CMKicker"]), Paragraph("Customer data quality is the main delivery risk", styles["CMTitle"]), Paragraph("Commercial success depends on repeatable onboarding and customer trust in the evidence model.", styles["CMSub"])])
    rows = [["Cost / risk", "Why it matters", "Mitigation hypothesis"], ["Dependency data collection", "Relationships may be scattered, stale, or contested.", "Customer-specific data mapping, owner workshops, and validation status."], ["Integration effort", "Identifiers, plans, and access controls vary by customer.", "Start with file imports; price implementation separately."], ["Sensitive data", "Infrastructure context may be restricted or security-sensitive.", "Customer-hosted/on-prem options; minimize retained document text."], ["AI trust", "Extraction and summaries can be wrong or misleading.", "Citations, review-only proposals, local inference, owner sign-off."], ["Incumbent overlap", "Existing platforms may already map services and dependencies.", "Compare workflow fit, exportability, deployment, and evidence control."], ["Long procurement", "Government and utilities can have lengthy approvals.", "Pilot with a bounded, synthetic or approved data subset."]]
    story.append(table_data(rows, styles, [36 * mm, 69 * mm, 52 * mm]))
    story.append(Spacer(1, 8))
    story.append(Paragraph("Go-to-market sequence", styles["CMHeading"]))
    for bullet in ["0–3 months: 8–12 discovery conversations, data-boundary review, and one safe exercise use case.", "3–6 months: customer-controlled pilot with named owners, one critical service cohort, and baseline preparation metrics.", "6–12 months: harden identity and audit, prove implementation repeatability, and compare against incumbent tools."]:
        story.append(Paragraph("• " + escape(bullet), styles["CMBullet"]))
    story.append(Paragraph("The key commercial gate is whether customers can produce a current, evidence-backed graph at a repeatable cost and whether the workflow offers value beyond existing continuity and event-management platforms. See docs/SOURCES.md for research and product sources.", styles["CMBody"]))
    document = SimpleDocTemplate(str(ROOT / "ContinuityMap_Business_Case.pdf"), pagesize=A4, rightMargin=17 * mm, leftMargin=17 * mm, topMargin=17 * mm, bottomMargin=18 * mm, title="ContinuityMap Business Case")
    document.build(story, onFirstPage=pdf_header_footer, onLaterPages=pdf_header_footer)


def build_evaluation_pdf() -> None:
    styles = base_styles()
    results = json.loads((ROOT / "data" / "evaluation_results.json").read_text())
    story = [Paragraph("SYNTHETIC EVALUATION · SNAPSHOT", styles["CMKicker"]), Paragraph("The software matches its test examples; real-world accuracy is untested", styles["CMTitle"]), Paragraph("This report checks a small invented network and hand-checked examples. It does not estimate real outages or prove that AI can understand customer plans.", styles["CMSub"]), Paragraph(results["scope"], styles["CMBody"])]
    graph = results["graph"]
    story.append(Paragraph("Dataset and integrity", styles["CMHeading"]))
    rows = [["Measure", "Result"], ["Services, facilities and resources / relationship records", f"{graph['nodes']} / {graph['edges']}"], ["Human-verified relationships", str(graph["verified_edges"])], ["Data checks", f"Passed={graph['integrity_valid']} · errors={graph['integrity_errors']} · warnings={graph['integrity_warnings']}"]]
    story.append(table_data(rows, styles, [64 * mm, 91 * mm]))
    extraction = results["dependency_extraction"]
    story.append(Paragraph("Simple text extraction check", styles["CMHeading"]))
    story.append(Paragraph("This check asks whether the app found the expected relationship in short fictional sentences. A correct suggestion matches the hand-checked answer; a false suggestion does not; a missed link was expected but not found.", styles["CMBody"]))
    rows = [["Correct", "False", "Missed", "Precision", "Recall", "F1", "Examples"], [str(extraction["tp"]), str(extraction["fp"]), str(extraction["fn"]), f"{extraction['precision']:.0%}", f"{extraction['recall']:.0%}", f"{extraction['f1']:.0%}", str(extraction["n_sentences"])]]
    story.append(table_data(rows, styles, [20 * mm, 20 * mm, 20 * mm, 24 * mm, 24 * mm, 20 * mm, 27 * mm]))
    story.append(Paragraph(extraction["method"] + " The test uses five short, clear, fictional sentences. Precision asks what share of suggestions were correct; recall asks what share of expected links were found; F1 combines those two scores. This is a parser check, not representative model evaluation.", styles["CMSmall"]))
    review = results["human_review_workflow"]
    story.append(Paragraph("Human-review state workflow", styles["CMHeading"]))
    story.append(Paragraph(f"{review['state_transition_correct']} of {review['cases']} synthetic status-transition cases matched their expected official inclusion state ({review['accuracy']:.0%}). {review['interpretation']}", styles["CMBody"]))
    unsupported = results["unsupported_edge_rate"]
    story.append(Paragraph("Evidence details on verified relationships", styles["CMHeading"]))
    story.append(Paragraph(f"{unsupported['verified_edges_missing_required_provenance']} of {graph['verified_edges']} verified fixture edges were missing required metadata ({unsupported['rate']:.0%}). {unsupported['interpretation']}", styles["CMBody"]))
    story.append(PageBreak())
    story.extend([Paragraph("FIXED SCENARIOS · UNCERTAINTY", styles["CMKicker"]), Paragraph("All three fictional test answers match", styles["CMTitle"]), Paragraph("The expected service lists were written down separately, then compared with the app output after each run.", styles["CMSub"])])
    rows = [["Scenario", "Expected", "Calculated", "Exact match?", "Precision", "Recall", "Time (ms)"]]
    for item in results["scenario_propagation"]["fixtures"]:
        rows.append([item["id"], str(item["expected"]), str(item["calculated"]), "Yes" if item["exact_match"] else "No", f"{item['precision']:.0%}", f"{item['recall']:.0%}", str(item["latency_ms"])])
    story.append(table_data(rows, styles, [48 * mm, 18 * mm, 20 * mm, 14 * mm, 18 * mm, 18 * mm, 18 * mm]))
    story.append(Spacer(1, 7))
    story.append(Paragraph(f"The app matched {results['scenario_propagation']['exact_fixture_matches']} of {results['scenario_propagation']['fixture_count']} expected lists exactly. Expected and calculated counts include the explicitly disrupted starting item. Precision is the share of app-listed items that also appear in the hand-checked list; recall is the share of expected items the app listed. F1 combines precision and recall into one score. Time is shown in milliseconds (thousandths of a second) and is only an informational local measurement.", styles["CMBody"]))
    story.append(Paragraph("Redundancy, cycle, and simultaneous disruption checks", styles["CMHeading"]))
    for bullet in ["ANY group: dependent remains available while one provider is available; it fails only when all declared alternatives are unavailable.", "ALL group: dependent becomes unavailable when any required provider is unavailable.", "Cycle example: the validator warns about a loop; repeated checks stop when no new services are added, and a disconnected group stays unaffected.", "Unknown/proposed paths: review nodes are returned separately; no unverified edge enters official status propagation."]:
        story.append(Paragraph("• " + escape(bullet), styles["CMBullet"]))
    uncertain = results["uncertainty"]
    rows = [["Scenario", "Open review nodes"]] + [[item["scenario"], ", ".join(item["review_nodes"]) or "None"] for item in uncertain["scenario_reviews"]]
    story.append(Paragraph("Uncertainty visibility", styles["CMHeading"]))
    story.append(table_data(rows, styles, [58 * mm, 97 * mm]))
    story.append(PageBreak())
    story.extend([Paragraph("LATENCY · EVIDENCE · LIMITATIONS", styles["CMKicker"]), Paragraph("A useful software test, not a validated resilience model", styles["CMTitle"]), Paragraph("The current benchmark verifies code behavior and evidence handling within a small, fixed fixture.", styles["CMSub"])])
    latency = results["latency"]
    rows = [["Latency measure", "Observed"], ["Repeated deterministic runs", str(latency["n"])], ["Median", f"{latency['median_ms']} ms"], ["P95", f"{latency['p95_ms']} ms"], ["Environment", latency["environment"]]]
    story.append(table_data(rows, styles, [52 * mm, 103 * mm]))
    evidence = results["evidence_summary"]
    story.append(Paragraph("Do evidence references resolve?", styles["CMHeading"]))
    story.append(Paragraph(f"{evidence['checked_refs']} evidence reference identifiers were checked and all resolve to graph records ({evidence['evidence_ref_validity']:.0%}). {evidence['interpretation']}", styles["CMBody"]))
    story.append(Paragraph("PDF report completeness", styles["CMHeading"]))
    for name, complete in results["report_completeness"].items():
        story.append(Paragraph(f"• {name.title()} report required sections present: {'Yes' if complete else 'No'}", styles["CMBullet"]))
    story.append(Paragraph("Evaluation limits", styles["CMHeading"]))
    for limitation in results["limitations"]:
        story.append(Paragraph("• " + escape(limitation), styles["CMBullet"]))
    story.append(Paragraph("Next evaluation steps", styles["CMHeading"]))
    for item in ["Create a customer-approved, independently adjudicated relationship set with ambiguous negatives and contradictory statements.", "Measure owner confirmation precision/recall and source-grounded quote quality, not only workflow state transitions.", "Benchmark graph size and service-layer latency on representative but non-sensitive customer-approved data.", "Conduct a structured review with service owners to test whether paths and proposed actions are understandable and useful."]:
        story.append(Paragraph("• " + escape(item), styles["CMBullet"]))
    document = SimpleDocTemplate(str(ROOT / "ContinuityMap_Evaluation_Report.pdf"), pagesize=A4, rightMargin=17 * mm, leftMargin=17 * mm, topMargin=17 * mm, bottomMargin=18 * mm, title="ContinuityMap Evaluation Report")
    document.build(story, onFirstPage=pdf_header_footer, onLaterPages=pdf_header_footer)


def wrap_draw(canvas, text, x, y_top, width, font="Helvetica", size=9, color=INK, leading=None, max_height=1000):
    leading = leading or size * 1.3
    style = ParagraphStyle("temp", fontName=font, fontSize=size, leading=leading, textColor=color, spaceAfter=0)
    escaped = escape(str(text)).replace("&lt;br/&gt;", "<br/>")
    paragraph = Paragraph(escaped, style)
    _, height = paragraph.wrap(width, max_height)
    paragraph.drawOn(canvas, x, y_top - height)
    return height


def build_pitch_pdf() -> None:
    deck = json.loads((ROOT / "data" / "deck_content.json").read_text())
    page_size = (13.333 * inch, 7.5 * inch)
    output = ROOT / "ContinuityMap_Pitch_Deck.pdf"
    canvas = __import__("reportlab.pdfgen.canvas", fromlist=["Canvas"]).Canvas(str(output), pagesize=page_size, pageCompression=1)
    w, h = page_size
    items = deck["slides"]
    for index, item in enumerate(items):
        if index == 0:
            canvas.setFillColor(NAVY)
            canvas.rect(0, 0, w, h, fill=1, stroke=0)
            canvas.setFillColor(TEAL)
            canvas.rect(0, 0, 0.13 * inch, h, fill=1, stroke=0)
            canvas.setFillColor(colors.HexColor("#70C8C4")); canvas.setFont("Helvetica-Bold", 10)
            canvas.drawString(0.75 * inch, 6.8 * inch, "CONTINUITYMAP")
            wrap_draw(canvas, "Essential services<br/>rely on one another.", 0.75 * inch, 5.98 * inch, 8.7 * inch, "Helvetica-Bold", 30, colors.white, 35)
            wrap_draw(canvas, deck["subtitle"], 0.77 * inch, 4.12 * inch, 8.2 * inch, "Helvetica", 15, colors.HexColor("#CDDCE2"), 19)
            wrap_draw(canvas, item["subtitle"], 0.77 * inch, 3.55 * inch, 8.1 * inch, "Helvetica", 11, colors.HexColor("#AAC0CA"), 15)
            cards = item["visualItems"]
            bx0 = 0.75 * inch
            for i, row in enumerate(cards):
                bx = bx0 + i * 3.0 * inch
                canvas.setFillColor(colors.HexColor("#1C4050" if i == 0 else "#193747")); canvas.setStrokeColor(colors.HexColor("#355966")); canvas.roundRect(bx, 1.42 * inch, 2.55 * inch, 1.04 * inch, 9, fill=1, stroke=1)
                canvas.setFillColor(colors.HexColor("#76C6C1")); canvas.setFont("Helvetica-Bold", 8); canvas.drawString(bx + 0.16 * inch, 2.16 * inch, row["title"].upper())
                wrap_draw(canvas, row["text"], bx + 0.16 * inch, 1.98 * inch, 2.18 * inch, "Helvetica-Bold", 9, colors.white, 11)
                if i < len(cards) - 1:
                    canvas.setFillColor(colors.HexColor("#79BFBA")); canvas.setFont("Helvetica-Bold", 18); canvas.drawCentredString(bx + 2.78 * inch, 1.89 * inch, "→")
            canvas.setFillColor(colors.HexColor("#213F4D")); canvas.roundRect(0.78 * inch, 0.6 * inch, 4.03 * inch, 0.36 * inch, 6, fill=1, stroke=0)
            canvas.setFillColor(colors.HexColor("#C5D5DB")); canvas.setFont("Helvetica-Bold", 7); canvas.drawCentredString(2.79 * inch, 0.73 * inch, "FICTIONAL DEMONSTRATION · PLANNING SUPPORT ONLY")
            canvas.setFillColor(colors.HexColor("#91A8B0")); canvas.setFont("Helvetica", 8); canvas.drawRightString(w - 0.6 * inch, 0.34 * inch, "1")
            canvas.showPage()
            continue

        light = colors.HexColor("#F3F6F7")
        canvas.setFillColor(light); canvas.rect(0, 0, w, h, fill=1, stroke=0)
        canvas.setFillColor(TEAL); canvas.rect(0, 0, 0.1 * inch, h, fill=1, stroke=0)
        accent = {"amber": AMBER, "coral": CORAL, "blue": colors.HexColor("#5E80AA")}.get(item["accent"], TEAL)
        pale = {"amber": colors.HexColor("#FFF4DC"), "coral": colors.HexColor("#FFF0EB"), "blue": colors.HexColor("#EDF2F8")}.get(item["accent"], colors.HexColor("#E5F4F2"))
        canvas.setFillColor(accent); canvas.setFont("Helvetica-Bold", 8); canvas.drawString(0.62 * inch, 7.08 * inch, item["section"].upper())
        wrap_draw(canvas, item["title"], 0.62 * inch, 6.82 * inch, 11.9 * inch, "Helvetica-Bold", 22, NAVY, 26)
        wrap_draw(canvas, item["subtitle"], 0.64 * inch, 6.18 * inch, 11.9 * inch, "Helvetica", 10.5, MUTED, 13)
        canvas.setStrokeColor(LINE); canvas.line(0.62 * inch, 5.72 * inch, 12.72 * inch, 5.72 * inch)

        canvas.setFillColor(colors.HexColor("#91A0A6")); canvas.setFont("Helvetica-Bold", 7.5); canvas.drawString(0.68 * inch, 5.42 * inch, "PLANNING INSIGHT")
        y = 5.05 * inch
        for bullet in item["bullets"]:
            canvas.setFillColor(accent); canvas.circle(0.75 * inch, y - 0.08 * inch, 0.045 * inch, fill=1, stroke=0)
            height = wrap_draw(canvas, bullet, 0.92 * inch, y, 4.48 * inch, "Helvetica", 10.2, INK, 13, 0.75 * inch)
            y -= max(0.77 * inch, height + 0.32 * inch)
        canvas.setFillColor(pale); canvas.setStrokeColor(pale); canvas.roundRect(0.67 * inch, 0.6 * inch, 4.88 * inch, 0.8 * inch, 7, fill=1, stroke=1)
        canvas.setFillColor(accent); canvas.setFont("Helvetica-Bold", 7); canvas.drawString(0.86 * inch, 1.18 * inch, "TAKEAWAY")
        wrap_draw(canvas, item["takeaway"], 0.86 * inch, 1.04 * inch, 4.48 * inch, "Helvetica-Bold", 8.6, INK, 11)

        x, panel_y, panel_w, panel_h = 5.92 * inch, 1.04 * inch, 6.75 * inch, 4.48 * inch
        canvas.setFillColor(colors.white); canvas.setStrokeColor(LINE); canvas.roundRect(x, panel_y, panel_w, panel_h, 8, fill=1, stroke=1)
        wrap_draw(canvas, item["visualTitle"], x + 0.3 * inch, panel_y + panel_h - 0.22 * inch, 6.1 * inch, "Helvetica-Bold", 12, NAVY, 15)
        draw_pdf_visual(canvas, item, x + 0.25 * inch, panel_y + 0.23 * inch, panel_w - 0.5 * inch, panel_h - 0.9 * inch, accent, pale)
        canvas.setStrokeColor(LINE); canvas.line(0.55 * inch, 0.46 * inch, 12.78 * inch, 0.46 * inch)
        canvas.setFillColor(MUTED); canvas.setFont("Helvetica", 7); canvas.drawString(0.6 * inch, 0.27 * inch, "CONTINUITYMAP · SYNTHETIC PLANNING DEMONSTRATION")
        canvas.drawRightString(12.7 * inch, 0.27 * inch, str(index + 1))
        canvas.showPage()
    canvas.save()


def draw_pdf_visual(canvas, item, x, y, w, h, accent, pale):
    rows = item["visualItems"]
    if item["visualStyle"] == "flow":
        gap = 0.12 * inch
        box_w = (w - 3 * gap - 0.3 * inch) / 4
        top = y + h / 2 - 0.64 * inch
        for i, row in enumerate(rows):
            bx = x + i * (box_w + gap + 0.1 * inch)
            canvas.setFillColor(pale if i == 0 else PALE); canvas.setStrokeColor(LINE); canvas.roundRect(bx, top, box_w, 1.28 * inch, 7, fill=1, stroke=1)
            canvas.setFillColor(accent); canvas.circle(bx + 0.22 * inch, top + 1.05 * inch, 0.12 * inch, fill=1, stroke=0)
            wrap_draw(canvas, row["title"], bx + 0.14 * inch, top + 0.85 * inch, box_w - 0.22 * inch, "Helvetica-Bold", 8.6, NAVY, 10)
            wrap_draw(canvas, row["text"], bx + 0.14 * inch, top + 0.52 * inch, box_w - 0.22 * inch, "Helvetica", 6.8, MUTED, 8)
            if i < len(rows) - 1:
                canvas.setFillColor(accent); canvas.setFont("Helvetica-Bold", 16); canvas.drawCentredString(bx + box_w + gap / 2 + 0.03 * inch, top + 0.55 * inch, "›")
        return
    if item["visualStyle"] == "record":
        canvas.setFillColor(pale); canvas.setStrokeColor(colors.HexColor("#D7EBE7")); canvas.roundRect(x + 0.25 * inch, y + 0.48 * inch, w - 0.5 * inch, h - 0.85 * inch, 8, fill=1, stroke=1)
        wrap_draw(canvas, rows[0]["title"], x + 0.5 * inch, y + h - 0.55 * inch, w - 1.0 * inch, "Helvetica-Bold", 14, NAVY, 18)
        wrap_draw(canvas, rows[0]["text"], x + 0.5 * inch, y + h - 1.25 * inch, w - 1.0 * inch, "Helvetica", 9, INK, 13)
        return
    if item["visualStyle"] == "logic":
        card_w = (w - 0.3 * inch) / 2
        for i, row in enumerate(rows):
            bx = x + i * (card_w + 0.3 * inch)
            fill = colors.HexColor("#F1F8F5") if i == 0 else colors.HexColor("#F8F8F6")
            canvas.setFillColor(fill); canvas.setStrokeColor(LINE); canvas.roundRect(bx, y + 0.4 * inch, card_w, h - 0.7 * inch, 8, fill=1, stroke=1)
            title, *rest = row["title"].split("·")
            canvas.setFillColor(TEAL if i == 0 else AMBER); canvas.setFont("Helvetica-Bold", 14); canvas.drawString(bx + 0.2 * inch, y + h - 0.65 * inch, title.strip())
            if rest:
                canvas.setFillColor(MUTED); canvas.setFont("Helvetica-Bold", 8); canvas.drawRightString(bx + card_w - 0.2 * inch, y + h - 0.62 * inch, rest[0].strip())
            wrap_draw(canvas, row["text"], bx + 0.2 * inch, y + h - 1.0 * inch, card_w - 0.4 * inch, "Helvetica", 8.5, INK, 12)
        return
    if item["visualStyle"] == "trace":
        top = y + h - 0.35 * inch
        for i, row in enumerate(rows):
            cy = top - i * 0.75 * inch
            color = AMBER if i == 3 else CORAL if i == 2 else accent
            canvas.setFillColor(colors.white); canvas.setStrokeColor(color); canvas.circle(x + 0.15 * inch, cy - 0.2 * inch, 0.09 * inch, fill=1, stroke=1)
            if i < len(rows) - 1:
                canvas.setStrokeColor(color); canvas.setLineWidth(1.2); canvas.line(x + 0.15 * inch, cy - 0.3 * inch, x + 0.15 * inch, cy - 0.75 * inch)
            fill = colors.HexColor("#FFF5DF") if i == 3 else PALE
            canvas.setFillColor(fill); canvas.setStrokeColor(LINE); canvas.roundRect(x + 0.42 * inch, cy - 0.49 * inch, w - 0.55 * inch, 0.57 * inch, 5, fill=1, stroke=1)
            canvas.setFillColor(color); canvas.setFont("Helvetica-Bold", 7); canvas.drawString(x + 0.62 * inch, cy - 0.14 * inch, row["title"].upper())
            wrap_draw(canvas, row["text"], x + 2.0 * inch, cy - 0.07 * inch, w - 2.34 * inch, "Helvetica", 8, INK, 9)
        return
    if item["visualStyle"] == "compare":
        for i, row in enumerate(rows):
            by = y + h - (i + 1) * 1.33 * inch
            fill = colors.HexColor("#E5F4F2") if i == 0 else colors.HexColor("#EDF2F8")
            canvas.setFillColor(fill); canvas.setStrokeColor(LINE); canvas.roundRect(x + 0.15 * inch, by, w - 0.3 * inch, 1.08 * inch, 7, fill=1, stroke=1)
            canvas.setFillColor(TEAL if i == 0 else colors.HexColor("#5E80AA")); canvas.setFont("Helvetica-Bold", 9); canvas.drawString(x + 0.38 * inch, by + 0.77 * inch, row["title"])
            wrap_draw(canvas, row["text"], x + 0.38 * inch, by + 0.64 * inch, w - 0.8 * inch, "Helvetica", 8.3, INK, 11)
        return
    if item["visualStyle"] == "revenue":
        canvas.setFillColor(MUTED); canvas.setFont("Helvetica-Bold", 7)
        canvas.drawString(x + 0.12 * inch, y + h - 0.4 * inch, "ORGANIZATIONS")
        canvas.drawString(x + 2.6 * inch, y + h - 0.4 * inch, "YEAR-ONE REVENUE / GROSS PROFIT")
        for i, row in enumerate(rows):
            cy = y + h - (i + 1) * 0.68 * inch - 0.5 * inch
            canvas.setFillColor(colors.HexColor("#F2F7F7") if i % 2 == 0 else colors.HexColor("#F8FAFA")); canvas.roundRect(x + 0.05 * inch, cy, w - 0.1 * inch, 0.52 * inch, 4, fill=1, stroke=0)
            canvas.setFillColor(NAVY); canvas.setFont("Helvetica-Bold", 8.5); canvas.drawString(x + 0.23 * inch, cy + 0.18 * inch, row["title"])
            canvas.setFillColor(TEAL); canvas.drawString(x + 2.0 * inch, cy + 0.18 * inch, row["text"])
        wrap_draw(canvas, "Model: $92k recurring + $28k onboarding · $30k direct year-one cost per organization.", x + 0.14 * inch, y + 0.23 * inch, w - 0.2 * inch, "Helvetica", 6.8, MUTED, 8)
        return
    if item["visualStyle"] == "timeline":
        base = y + 1.2 * inch
        canvas.setStrokeColor(colors.HexColor("#CBDADD")); canvas.setLineWidth(2); canvas.line(x + 0.3 * inch, base, x + w - 0.3 * inch, base)
        for i, row in enumerate(rows):
            bx = x + i * 2.08 * inch
            canvas.setFillColor(colors.white); canvas.setStrokeColor(accent); canvas.circle(bx + 0.3 * inch, base, 0.13 * inch, fill=1, stroke=1)
            canvas.setFillColor(pale if i == 0 else PALE); canvas.setStrokeColor(LINE); canvas.roundRect(bx, y + 0.05 * inch, 1.92 * inch, 1.3 * inch, 6, fill=1, stroke=1)
            canvas.setFillColor(accent); canvas.setFont("Helvetica-Bold", 9); canvas.drawString(bx + 0.15 * inch, y + 1.12 * inch, row["title"])
            wrap_draw(canvas, row["text"], bx + 0.15 * inch, y + 0.96 * inch, 1.62 * inch, "Helvetica", 7, INK, 9)
        return
    cols = 2 if len(rows) > 3 else len(rows)
    gap = 0.28 * inch
    card_w = (w - gap * (cols - 1)) / cols
    card_h = 1.18 * inch if len(rows) > 3 else 1.65 * inch
    total_rows = (len(rows) + cols - 1) // cols
    for i, row in enumerate(rows):
        row_index, col = divmod(i, cols)
        bx = x + col * (card_w + gap)
        by = y + h - 0.35 * inch - (row_index + 1) * card_h - row_index * 0.25 * inch
        canvas.setFillColor(pale if i == 0 else PALE); canvas.setStrokeColor(LINE); canvas.roundRect(bx, by, card_w, card_h, 7, fill=1, stroke=1)
        canvas.setFillColor(accent); canvas.rect(bx, by, 0.05 * inch, card_h, fill=1, stroke=0)
        wrap_draw(canvas, row["title"], bx + 0.18 * inch, by + card_h - 0.17 * inch, card_w - 0.34 * inch, "Helvetica-Bold", 8.8, NAVY, 10)
        wrap_draw(canvas, row["text"], bx + 0.18 * inch, by + card_h - 0.55 * inch, card_w - 0.34 * inch, "Helvetica", 7.8, MUTED, 10)


def main() -> None:
    build_idea_docx()
    build_idea_pdf()
    build_business_pdf()
    build_evaluation_pdf()
    build_pitch_pdf()
    print("Wrote idea DOCX/PDF, business PDF, evaluation PDF, and 14-slide deck PDF.")


if __name__ == "__main__":
    main()
